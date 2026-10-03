#!/usr/bin/env python3
"""Genera el borrador DOCX del reporte de estadía (R-ADC-08-07).

Construye el documento a partir de la plantilla institucional (que se recibe
como argumento y no se versiona aquí) y de las fuentes Markdown consolidadas en
``docs/08-reporte-estadia/``. No inventa contenido: los capítulos se renderizan
desde los Markdown; el Resumen, la tabla de fases y la estructura de anexos son
curaduría explícita del guion de entrega documentado en este mismo script.

Uso:
    python scripts/report/generate_docx.py --template RUTA/plantilla.docx \
        [--output docs/08-reporte-estadia/R-ADC-08-07_Reporte_Estadia_Borrador_V1.docx]

Requisitos: python-docx (``python -m pip install python-docx``).
Las figuras renderizadas y capturas sanitizadas deben existir en
``docs/08-reporte-estadia/anexos/`` (ver scripts/report/README.md).
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

REPO_ROOT = Path(__file__).resolve().parents[2]
DOCS = REPO_ROOT / 'docs' / '08-reporte-estadia'
ANEXOS = DOCS / 'anexos'
CAPTURAS = ANEXOS / 'capturas'
FIGURAS = ANEXOS / 'figuras'

BOLD_RE = re.compile(r'\*\*(.+?)\*\*')
LINK_RE = re.compile(r'\[([^\]]+)\]\(([^)]+)\)')
ORDERED_RE = re.compile(r'^\d+\.\s+')

RESUMEN = (
    'El presente reporte describe el proyecto de modernización del control de inventarios de '
    'comercio exterior (Anexo 24) de PROCESADORA DE ALIMENTOS CALE, S.A. de C.V. El sistema de '
    'referencia concentra la operación en una base SQL Server con procesos almacenados y presenta '
    'riesgos documentados durante la auditoría: exposición de credenciales en interfaz, permisos '
    'administrativos amplios, reportes que fallan con error genérico y cargas de archivos sin '
    'diagnóstico detallado. El objetivo fue desarrollar una aplicación web segura y verificable '
    'que consulte y cargue información sobre la base operativa existente sin modificarla sin '
    'autorización. La solución implementada combina Angular y Spring Boot con arquitectura '
    'hexagonal, acceso exclusivo mediante procedimientos almacenados versionados (regla SP-FIRST), '
    'staging aislado con validación por archivo/fila/columna y huella anticuplicados, reportes '
    'consolidados con exportación, administración con permisos y bitácora. La metodología siguió un '
    'ciclo por fases: auditoría del sistema legacy, descubrimiento read-only de la base autoritativa, '
    'diseño de arquitectura, desarrollo por módulos con pruebas unitarias e integración sobre SQL '
    'Server efímero, auditoría funcional interactiva de extremo a extremo y verificación continua en '
    'integración (CI). Los resultados principales: 79 capacidades legacy clasificadas con evidencia '
    '(5 equivalentes, 26 rediseñadas, 12 parciales, 6 faltantes, 5 bloqueadas, 6 consolidadas y 19 '
    'desconocidas), once reportes implementados, cuatro flujos de carga con staging, confirmación '
    'autoritativa de pedimentos, 131 de 131 pruebas de frontend exitosas, sin hallazgos E2E de '
    'severidad alta abiertos, 77 de 77 objetos SQL requeridos presentes en el entorno LIVE y CI '
    'verde. Como limitaciones, permanecen decisiones de negocio pendientes (saldos, motor de '
    'descargos, confirmación de facturación, alcance del dashboard y reglas PED-005/006), doce '
    'capacidades esperando contrato o evidencia externa y la validación con usuarios reales; por '
    'ello este documento es un borrador académico y no declara el cierre total del proyecto.'
)

PHASE_TABLE = [
    ('1', 'Auditoría del sistema legacy', 'Riesgos y línea base funcional'),
    ('2', 'Descubrimiento de la base autoritativa', 'Regla SP-FIRST y límites de ejecución'),
    ('3', 'Diseño de arquitectura', 'Arquitectura hexagonal y dos orígenes de datos'),
    ('4', 'Catálogos', 'Consultas read-only de maestros y auxiliares'),
    ('5', 'Operaciones read-only', 'Consultas con periodo, filtros y paginación'),
    ('6', 'Administración y seguridad', 'Usuarios, perfiles, permisos y bitácora auditables'),
    ('7', 'Reportes', 'Once consolidados con paginación y exportación XLSX'),
    ('8', 'Staging de pedimentos', 'Parser, validación, hash y preview aislados'),
    ('9', 'Confirmación autoritativa de pedimentos', 'Transacción, idempotencia y rollback verificados'),
    ('10', 'Staging de materiales y productos', 'Importación validada sin efectos en catálogo'),
    ('11', 'Staging de facturación', 'Plantilla activa, validación y duplicados controlados'),
    ('12', 'Paridad funcional', 'Matriz de 79 capacidades clasificada con evidencia'),
    ('13', 'Auditoría E2E', 'Sweep interactivo de la aplicación completa'),
    ('14', 'Correcciones', 'Hallazgos P1/P2 cerrados con pruebas de regresión'),
    ('15', 'Cierre técnico V1', 'Snapshot consolidado y punto de referencia RC1'),
]

ANNEX_B_INTERFACES = [
    ('01-login.png', 'Anexo B.1. Inicio de sesión de la aplicación.'),
    ('02-dashboard.png', 'Anexo B.2. Panel principal con accesos filtrados por permiso.'),
    ('03-materiales.png', 'Anexo B.3. Consulta del catálogo de materiales.'),
    ('04-productos.png', 'Anexo B.4. Consulta del catálogo de productos.'),
    ('05-estructuras.png', 'Anexo B.5. Consulta de estructuras (sin datos en el entorno).'),
    ('06-catalogos.png', 'Anexo B.6. Catálogos auxiliares.'),
    ('07-datos-generales.png', 'Anexo B.7. Ficha empresarial (sin registro en el entorno).'),
    ('08-socios-comerciales.png', 'Anexo B.8. Socios comerciales.'),
]

ANNEX_C_OPERACIONES = [
    ('10-entradas.png', 'Anexo C.1. Consulta de entradas.'),
    ('11-salidas.png', 'Anexo C.2. Consulta de salidas.'),
    ('12-materiales-utilizados.png', 'Anexo C.3. Consulta de materiales utilizados.'),
    ('13-activos-fijos.png', 'Anexo C.4. Consulta de activos fijos.'),
    ('17-usuarios.png', 'Anexo C.5. Administración de usuarios.'),
    ('18-perfiles.png', 'Anexo C.6. Administración de perfiles.'),
]

ANNEX_D_CARGAS = [
    ('09a-importaciones-inicial.png', 'Anexo D.1. Importación de catálogos: selección de archivo.'),
    ('09b-importaciones-preview.png', 'Anexo D.2. Importación de catálogos: previsualización validada.'),
    ('14a-pedimentos-inicial.png', 'Anexo D.3. Carga de pedimentos: estado inicial.'),
    ('14b-pedimentos-validacion.png', 'Anexo D.4. Carga de pedimentos: validación con errores estructurados.'),
    ('16a-facturacion-inicial.png', 'Anexo D.5. Carga de facturación: estado inicial.'),
    ('16b-facturacion-validada.png', 'Anexo D.6. Carga de facturación: archivo validado.'),
]

ANNEX_E_REPORTES = [
    ('15-reportes.png', 'Anexo E.1. Reportes consolidados (opción F4 sin datos).'),
    ('19-bitacora.png', 'Anexo E.2. Bitácora de eventos con filtros.'),
]

FIXTURES_TABLE = [
    ('fixture-pedimento-sintetico.xlsx', 'Pedimento sintético (produce el error esperado PED-003)'),
    ('fixture-materiales.xlsx', 'Material sintético para importación de catálogo'),
    ('fixture-productos.xlsx', 'Producto sintético para importación de catálogo'),
    ('fixture-facturacion.xlsx', 'Facturación sintética (hoja FACTURAS)'),
    ('fixture-invalido-validacion.xlsx', 'Archivo inválido para evidencia de errores de estructura'),
]


# ---------------------------------------------------------------------------
# Utilidades de documento
# ---------------------------------------------------------------------------

def set_paragraph_text(paragraph, text: str) -> None:
    """Reemplaza el texto de un párrafo conservando el formato del primer run."""
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ''


def add_runs_with_bold(paragraph, text: str) -> None:
    """Agrega runs aplicando negritas ``**`` y limpiando enlaces/backticks."""
    text = LINK_RE.sub(r'\1', text)
    text = text.replace('`', '')
    position = 0
    for match in BOLD_RE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        run = paragraph.add_run(match.group(1))
        run.bold = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style='List Paragraph')
    paragraph.paragraph_format.left_indent = Cm(0.75)
    paragraph.add_run('• ')
    add_runs_with_bold(paragraph, text)


def add_numbered(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style='List Paragraph')
    paragraph.paragraph_format.left_indent = Cm(0.75)
    add_runs_with_bold(paragraph, text)


def add_bold_line(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True


def add_italic_note(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    source = doc.add_paragraph()
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source_run = source.add_run('Fuente: Elaboración propia.')
    source_run.font.size = Pt(8)


def add_figure(doc: Document, image_path: Path, caption: str, width_cm: float = 15.5) -> None:
    if not image_path.exists():
        raise FileNotFoundError(f'Figura no encontrada: {image_path}')
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))
    add_caption(doc, caption)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), '8C8C8C')
        borders.append(element)
    tbl_pr.append(borders)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], font_size: int = 9) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = True
    set_table_borders(table)
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        paragraph = header_cells[index].paragraphs[0]
        run = paragraph.add_run(header)
        run.bold = True
        run.font.size = Pt(font_size)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            paragraph = cells[index].paragraphs[0]
            run = paragraph.add_run(str(value))
            run.font.size = Pt(font_size)
    doc.add_paragraph()


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def add_chapter(doc: Document, title: str, first: bool = False) -> None:
    if not first:
        add_page_break(doc)
    doc.add_paragraph(title, style='Heading 1')


def add_subheading(doc: Document, title: str) -> None:
    doc.add_paragraph(title, style='Heading 2')


def add_inner_heading(doc: Document, title: str) -> None:
    add_bold_line(doc, title)


# ---------------------------------------------------------------------------
# Render de Markdown
# ---------------------------------------------------------------------------

def _flush_paragraph(doc: Document, buffer: list[str]) -> None:
    if not buffer:
        return
    text = ' '.join(part.strip() for part in buffer if part.strip())
    if text:
        paragraph = doc.add_paragraph()
        add_runs_with_bold(paragraph, text)
    buffer.clear()


def _render_table(doc: Document, lines: list[str]) -> None:
    rows: list[list[str]] = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip('|').split('|')]
        rows.append(cells)
    if len(rows) >= 2 and set(rows[1][0]) <= {'-', ':'}:
        header, body = rows[0], rows[2:]
    else:
        header, body = rows[0], rows[1:]
    header = [cell for cell in header if cell != '']
    body = [row[:len(header)] for row in body]
    add_table(doc, header, body)


def render_markdown(doc: Document, text: str, skip_h1: bool = True, demote_h2: bool = False,
                    skip_titles: tuple[str, ...] = (), skip_blockquotes_with: tuple[str, ...] = ()) -> None:
    """Renderiza Markdown simple (headings, párrafos, listas, tablas, citas)."""
    lines = text.splitlines()
    buffer: list[str] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped or stripped == '---':
            _flush_paragraph(doc, buffer)
            index += 1
            continue
        if stripped.startswith('|') and index + 1 < len(lines) and lines[index + 1].strip().startswith('|'):
            _flush_paragraph(doc, buffer)
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                table_lines.append(lines[index])
                index += 1
            _render_table(doc, table_lines)
            continue
        if stripped.startswith('#'):
            _flush_paragraph(doc, buffer)
            level = len(stripped) - len(stripped.lstrip('#'))
            title = stripped[level:].strip()
            if any(title.startswith(skip) or skip in title for skip in skip_titles):
                index += 1
                continue
            if level == 1:
                if not skip_h1:
                    doc.add_paragraph(title, style='Heading 1')
            elif level == 2:
                if demote_h2:
                    add_bold_line(doc, title)
                else:
                    add_subheading(doc, title)
            else:
                add_inner_heading(doc, title)
            index += 1
            continue
        if stripped.startswith('> '):
            _flush_paragraph(doc, buffer)
            content = stripped[2:].strip()
            if not any(token in content for token in skip_blockquotes_with):
                add_italic_note(doc, content)
            index += 1
            continue
        if stripped.startswith('- '):
            _flush_paragraph(doc, buffer)
            while index < len(lines) and lines[index].strip().startswith('- '):
                add_bullet(doc, lines[index].strip()[2:].strip())
                index += 1
            continue
        if ORDERED_RE.match(stripped):
            _flush_paragraph(doc, buffer)
            while index < len(lines) and ORDERED_RE.match(lines[index].strip()):
                item = ORDERED_RE.sub('', lines[index].strip())
                add_numbered(doc, item)
                index += 1
            continue
        buffer.append(line)
        index += 1
    _flush_paragraph(doc, buffer)


def read_source(name: str) -> str:
    return (DOCS / name).read_text(encoding='utf-8')


def read_source_until(name: str, marker: str) -> str:
    text = read_source(name)
    position = text.find(marker)
    return text[:position] if position >= 0 else text


def split_sections(markdown_text: str) -> dict[str, str]:
    """Divide un Markdown por encabezados ``## `` en secciones por título."""
    sections: dict[str, str] = {}
    current_title: str | None = None
    current_lines: list[str] = []
    for line in markdown_text.splitlines():
        if line.startswith('## '):
            if current_title is not None:
                sections[current_title] = '\n'.join(current_lines)
            current_title = line[3:].strip()
            current_lines = []
        else:
            current_lines.append(line)
    if current_title is not None:
        sections[current_title] = '\n'.join(current_lines)
    return sections


def split_phases(markdown_text: str) -> dict[int, str]:
    """Divide desarrollo.md en fases numeradas ``## Fase N —``."""
    phases: dict[int, str] = {}
    current: int | None = None
    current_lines: list[str] = []
    for line in markdown_text.splitlines():
        match = re.match(r'## Fase (\d+) ', line)
        if match:
            if current is not None:
                phases[current] = '\n'.join(current_lines)
            current = int(match.group(1))
            current_lines = [line]
        elif current is not None:
            current_lines.append(line)
    if current is not None:
        phases[current] = '\n'.join(current_lines)
    return phases


# ---------------------------------------------------------------------------
# Preparación de la plantilla
# ---------------------------------------------------------------------------

def fill_cover(doc: Document) -> None:
    replacements = {
        'Especificar nombre del proyecto': 'Modernización del control de inventarios de comercio exterior (Anexo 24)',
        'Especificar la empresa': 'PROCESADORA DE ALIMENTOS CALE, S.A. de C.V.',
        'Especificar nombres de estudiantes': 'Pendiente de completar por el estudiante',
        'Asesor académico: especificar nombre': 'Asesor académico: pendiente de completar',
        'Asesor industrial: especificar nombre': 'Asesor industrial: pendiente de completar',
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in replacements:
            set_paragraph_text(paragraph, replacements[text])


def replace_toc_with_field(doc: Document) -> None:
    """Elimina el índice manual y coloca un campo TOC actualizable por Word."""
    toc_paragraphs = [paragraph for paragraph in doc.paragraphs if paragraph.style.name.startswith('toc ')]
    if not toc_paragraphs:
        return
    anchor = None
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == 'Contenido' and paragraph.style.name == 'Normal':
            anchor = paragraph
            break
    if anchor is None:
        anchor = toc_paragraphs[0]
    for paragraph in toc_paragraphs:
        paragraph._element.getparent().remove(paragraph._element)
    field_paragraph = OxmlElement('w:p')
    run = OxmlElement('w:r')
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = 'TOC \\o "1-2" \\h \\z \\u'
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    placeholder = OxmlElement('w:t')
    placeholder.text = 'Actualice el campo en Word (F9) para generar la tabla de contenido.'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run.append(begin)
    run.append(instruction)
    run.append(separate)
    run.append(placeholder)
    run.append(end)
    field_paragraph.append(run)
    anchor._element.addnext(field_paragraph)


def clear_body_after_toc(doc: Document) -> None:
    """Elimina los esqueletos vacíos del cuerpo preservando portada y hojas previas."""
    start_index = None
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.style.name == 'Heading 1' and paragraph.text.strip() == 'Resumen':
            start_index = index
            break
    if start_index is None:
        return
    for paragraph in doc.paragraphs[start_index:]:
        paragraph._element.getparent().remove(paragraph._element)


# ---------------------------------------------------------------------------
# Cuerpo del reporte
# ---------------------------------------------------------------------------

def build_resumen(doc: Document) -> None:
    add_chapter(doc, 'Resumen', first=True)
    paragraph = doc.add_paragraph()
    add_runs_with_bold(paragraph, RESUMEN)


def build_introduccion(doc: Document) -> None:
    add_chapter(doc, '1. Introducción')
    add_subheading(doc, '1.1 Antecedentes')
    render_markdown(doc, read_source('antecedentes.md'))
    add_subheading(doc, '1.2 Planteamiento del problema')
    render_markdown(doc, read_source('planteamiento-problema.md'))
    add_subheading(doc, '1.3 Objetivo general')
    objetivos = read_source_until('objetivos.md', '## Estado de cumplimiento')
    render_markdown(doc, objetivos, skip_titles=('Objetivo general', 'Objetivos específicos'))
    add_inner_heading(doc, '1.3.1 Objetivos específicos')
    add_subheading(doc, '1.4 Hipótesis o supuesto')
    render_markdown(doc, read_source('hipotesis.md'), demote_h2=True)
    add_subheading(doc, '1.5 Justificación')
    render_markdown(doc, read_source('justificacion.md'))
    add_subheading(doc, '1.6 Alcance del proyecto')
    render_markdown(doc, read_source('alcance.md'), demote_h2=True)


def build_marco(doc: Document) -> None:
    add_chapter(doc, '2. Marco teórico y de referencia')
    add_subheading(doc, '2.1 Marco teórico')
    render_markdown(doc, read_source('marco-teorico.md'), demote_h2=True)
    lead = doc.add_paragraph()
    add_runs_with_bold(lead, 'La arquitectura general implementada se resume en la Figura 1.')
    add_figure(doc, FIGURAS / 'arquitectura-general.png',
               'Figura 1. Arquitectura general de la solución: Angular, API segura en Spring Boot y '
               'adaptador controlado a la base operativa.', width_cm=15.5)
    add_subheading(doc, '2.2 Marco de referencia')
    render_markdown(doc, read_source('marco-referencia.md'))


def build_desarrollo(doc: Document) -> None:
    add_chapter(doc, '3. Desarrollo: Métodos y técnicas')
    intro = doc.add_paragraph()
    add_runs_with_bold(
        intro,
        'El proyecto se ejecutó en quince fases verificables. La Tabla 1 resume el enfoque y el '
        'resultado de cada una; los apartados siguientes desarrollan problema, método, decisión, '
        'implementación, validación y limitación.',
    )
    add_table(doc, ['Fase', 'Foco', 'Resultado principal'], PHASE_TABLE, font_size=9)
    add_caption(doc, 'Tabla 1. Fases del proyecto y resultado principal.')

    phases = split_phases(read_source('desarrollo.md'))

    add_subheading(doc, '3.1 Análisis y descubrimiento')
    for number in (1, 2):
        render_markdown(doc, phases.get(number, ''), demote_h2=True)

    add_subheading(doc, '3.2 Diseño de arquitectura')
    for number in (3,):
        render_markdown(doc, phases.get(number, ''), demote_h2=True)

    add_subheading(doc, '3.3 Consultas, catálogos y administración')
    for number in (4, 5, 6, 7):
        render_markdown(doc, phases.get(number, ''), demote_h2=True)

    add_subheading(doc, '3.4 Cargas, staging y confirmación autoritativa')
    for number in (8, 9, 10, 11):
        render_markdown(doc, phases.get(number, ''), demote_h2=True)
    lead = doc.add_paragraph()
    add_runs_with_bold(
        lead,
        'La validación de una carga de pedimentos con errores estructurados se muestra en la '
        'Figura 2; el pipeline de carga validada que adoptó el proyecto se resume en la Figura 3.',
    )
    add_figure(doc, CAPTURAS / '14b-pedimentos-validacion.png',
               'Figura 2. Validación de pedimentos en staging con errores estructurados por fila y columna.')
    add_figure(doc, FIGURAS / 'pipeline-carga-validada.png',
               'Figura 3. Pipeline de carga validada: archivo, validación, staging y confirmación sólo '
               'con contrato aprobado.')

    add_subheading(doc, '3.5 Paridad, auditoría E2E y correcciones')
    for number in (12, 13, 14):
        render_markdown(doc, phases.get(number, ''), demote_h2=True)

    add_subheading(doc, '3.6 Cierre técnico V1')
    for number in (15,):
        render_markdown(doc, phases.get(number, ''), demote_h2=True)


def build_resultados(doc: Document) -> None:
    add_chapter(doc, '4. Análisis de resultados')
    sections = split_sections(read_source('resultados.md'))

    def find(fragment: str) -> str:
        for title, content in sections.items():
            if fragment in title:
                return content
        return ''

    add_subheading(doc, '4.1 Cobertura de paridad legacy')
    render_markdown(doc, find('Cobertura de paridad'))
    add_subheading(doc, '4.2 Módulos modernos entregados')
    render_markdown(doc, find('Módulos modernos'))
    add_subheading(doc, '4.3 Datasets observados')
    render_markdown(doc, find('Datasets observados'))
    add_subheading(doc, '4.4 Pruebas y calidad')
    render_markdown(doc, find('Pruebas y calidad'))
    add_subheading(doc, '4.5 SQL y despliegue')
    render_markdown(doc, find('SQL y despliegue'))
    add_subheading(doc, '4.6 Capacidades autoritativas y límites')
    render_markdown(doc, find('Capacidades autoritativas'))
    render_markdown(doc, find('Límites explícitos'))
    add_subheading(doc, '4.7 Evidencia visual')
    lead = doc.add_paragraph()
    add_runs_with_bold(
        lead,
        'Las Figuras 4 a 7 muestran evidencia sanitizada de la aplicación en ejecución (dashboard, '
        'catálogo de materiales, reportes y bitácora).',
    )
    add_figure(doc, CAPTURAS / '02-dashboard.png',
               'Figura 4. Dashboard posterior al inicio de sesión con accesos filtrados por permiso.')
    add_figure(doc, CAPTURAS / '03-materiales.png',
               'Figura 5. Consulta del catálogo de materiales con datos sintéticos de demostración.')
    add_figure(doc, CAPTURAS / '15-reportes.png',
               'Figura 6. Reporte F4 (CTM/desperdicio) en estado vacío dentro de la superficie de reportes.')
    add_figure(doc, CAPTURAS / '19-bitacora.png',
               'Figura 7. Bitácora de eventos con filtros y trazabilidad.')


def build_conclusiones(doc: Document) -> None:
    add_chapter(doc, '5. Conclusiones y recomendaciones')
    add_subheading(doc, '5.1 Conclusiones')
    add_italic_note(
        doc,
        'Nota: conclusiones provisionales basadas en evidencia técnica. La conclusión final de alcance '
        'queda pendiente de las decisiones de negocio y de la validación con usuarios descritas en '
        'el capítulo de recomendaciones.',
    )
    render_markdown(doc, read_source('conclusiones.md'),
                    skip_blockquotes_with=('FINAL_BUSINESS_CONCLUSION_PENDING',))
    add_subheading(doc, '5.2 Recomendaciones')
    render_markdown(doc, read_source('recomendaciones.md'), demote_h2=True)


def build_glosario(doc: Document) -> None:
    add_chapter(doc, '6. Glosario')
    render_markdown(doc, read_source('glosario.md'))


def build_bibliografia(doc: Document) -> None:
    add_chapter(doc, '7. Bibliografía')
    render_markdown(doc, read_source('bibliografia.md'), demote_h2=True)


def build_anexos(doc: Document) -> None:
    add_chapter(doc, '8. Anexos')
    intro = doc.add_paragraph()
    add_runs_with_bold(
        intro,
        'Los anexos se presentan en el orden en que son citados en el cuerpo. Las capturas fueron '
        'tomadas en entorno local con datos sintéticos o estados vacíos; no contienen información '
        'empresarial sensible. Los archivos de datos sintéticos se listan en el Anexo F.',
    )

    add_subheading(doc, 'Anexo A — Arquitectura y diseño')
    add_figure(doc, (REPO_ROOT / 'docs' / '03-diseno' / 'ERv3.png'),
               'Anexo A.1. Modelo entidad-relación propuesto (no incluye las tablas complementarias del '
               'esquema de aplicación).', width_cm=15.5)
    note = doc.add_paragraph()
    add_runs_with_bold(
        note,
        'La matriz de paridad completa (79 capacidades), el detalle de la auditoría E2E y el cierre '
        'técnico están disponibles como evidencia versionada en el repositorio del proyecto '
        '(docs/01-requerimientos, docs/05-pruebas).',
    )

    add_subheading(doc, 'Anexo B — Interfaces principales')
    for file_name, caption in ANNEX_B_INTERFACES:
        add_figure(doc, CAPTURAS / file_name, caption)

    add_subheading(doc, 'Anexo C — Operaciones y administración')
    for file_name, caption in ANNEX_C_OPERACIONES:
        add_figure(doc, CAPTURAS / file_name, caption)

    add_subheading(doc, 'Anexo D — Cargas y validaciones')
    for file_name, caption in ANNEX_D_CARGAS:
        add_figure(doc, CAPTURAS / file_name, caption)

    add_subheading(doc, 'Anexo E — Reportes y bitácora')
    for file_name, caption in ANNEX_E_REPORTES:
        add_figure(doc, CAPTURAS / file_name, caption)

    add_subheading(doc, 'Anexo F — Evidencia de pruebas y fixtures sintéticos')
    paragraph = doc.add_paragraph()
    add_runs_with_bold(
        paragraph,
        'La aplicación se validó con pruebas unitarias y de integración, auditoría funcional '
        'interactiva de extremo a extremo y verificación continua en CI. Los archivos sintéticos '
        'utilizados para la evidencia de cargas se listan en la Tabla 2.',
    )
    add_table(doc, ['Archivo sintético', 'Propósito'], FIXTURES_TABLE, font_size=9)
    add_caption(doc, 'Tabla 2. Fixtures sintéticos utilizados como evidencia.')


# ---------------------------------------------------------------------------
# Principal
# ---------------------------------------------------------------------------

def generate(template_path: Path, output_path: Path) -> None:
    if not template_path.exists():
        raise FileNotFoundError(f'Plantilla no encontrada: {template_path}')
    doc = Document(str(template_path))
    fill_cover(doc)
    replace_toc_with_field(doc)
    clear_body_after_toc(doc)

    build_resumen(doc)
    build_introduccion(doc)
    build_marco(doc)
    build_desarrollo(doc)
    build_resultados(doc)
    build_conclusiones(doc)
    build_glosario(doc)
    build_bibliografia(doc)
    build_anexos(doc)

    properties = doc.core_properties
    properties.title = 'R-ADC-08-07 Reporte de estadía — Borrador V1'
    properties.subject = 'Modernización del control de inventarios de comercio exterior (Anexo 24)'
    properties.author = 'Proyecto Anexo 24 — borrador académico'
    properties.comments = (
        'ACADEMIC_DRAFT=YES; FINAL_BUSINESS_CONCLUSION_PENDING=YES; '
        'documento generado desde fuentes Markdown versionadas.'
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def main() -> None:
    parser = argparse.ArgumentParser(description='Genera el borrador DOCX del reporte de estadía.')
    parser.add_argument('--template', required=True, help='Ruta de la plantilla oficial R-ADC-08-07 (.docx)')
    parser.add_argument('--output', default=str(DOCS / 'R-ADC-08-07_Reporte_Estadia_Borrador_V1.docx'),
                        help='Ruta de salida del DOCX')
    args = parser.parse_args()
    generate(Path(args.template).resolve(), Path(args.output).resolve())
    print(f'DOCX_GENERATED|{args.output}')


if __name__ == '__main__':
    main()
